import os
import django
from decimal import Decimal
from datetime import date
from collections import defaultdict
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'catalog_project.settings')
django.setup()

from products.models import (
    Company, Users, Supplieruser, SupplierManager, SupplierExecutive, EndUser,
    Customer, Supplier, CustomerBill, SupplierBill, Products, SupplierProduct
)
from django.db import connection

def export_bills_to_excel(output_path="Customer_and_Supplier_Bills_Report.xlsx"):
    print("=" * 75)
    print("  GENERATING COMPLETE MULTI-SHEET ENTERPRISE EXCEL WORKBOOK")
    print("=" * 75)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # Remove default blank sheet

    # -------------------------------------------------------------
    # STYLES & PALETTES
    # -------------------------------------------------------------
    font_main_title = Font(name="Segoe UI", size=15, bold=True, color="FFFFFF")
    font_sub_title = Font(name="Segoe UI", size=10, italic=True, color="EAECEE")
    font_sec_heading = Font(name="Segoe UI", size=12, bold=True, color="1B365D")
    
    font_col_header = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    font_data = Font(name="Segoe UI", size=10, color="000000")
    font_data_bold = Font(name="Segoe UI", size=10, bold=True, color="000000")
    font_data_code = Font(name="Segoe UI", size=10, bold=True, color="003366")
    font_total_label = Font(name="Segoe UI", size=10, bold=True, color="1B365D")
    font_total_val = Font(name="Segoe UI", size=10, bold=True, color="1B365D")

    # Header Color Fills
    fill_banner_navy = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    fill_banner_dark = PatternFill(start_color="212F3D", end_color="212F3D", fill_type="solid")
    
    fill_hdr_navy = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    fill_hdr_blue = PatternFill(start_color="2E5B88", end_color="2E5B88", fill_type="solid")
    fill_hdr_teal = PatternFill(start_color="117A65", end_color="117A65", fill_type="solid")
    fill_hdr_purple = PatternFill(start_color="5B2C6F", end_color="5B2C6F", fill_type="solid")
    fill_hdr_darkgreen = PatternFill(start_color="1E8449", end_color="1E8449", fill_type="solid")
    fill_hdr_darkamber = PatternFill(start_color="935116", end_color="935116", fill_type="solid")
    fill_hdr_steel = PatternFill(start_color="34495E", end_color="34495E", fill_type="solid")

    # Row Fills
    fill_zebra_light = PatternFill(start_color="F8F9F9", end_color="F8F9F9", fill_type="solid")
    fill_total_row = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    fill_badge_green = PatternFill(start_color="D4EFDF", end_color="D4EFDF", fill_type="solid")
    fill_badge_red = PatternFill(start_color="FADBD8", end_color="FADBD8", fill_type="solid")
    fill_badge_blue = PatternFill(start_color="D6EAF8", end_color="D6EAF8", fill_type="solid")
    fill_badge_yellow = PatternFill(start_color="FCF3CF", end_color="FCF3CF", fill_type="solid")
    fill_sub_section = PatternFill(start_color="EAEDED", end_color="EAEDED", fill_type="solid")

    # Borders
    thin_border_line = Side(style='thin', color='D5D8DC')
    medium_top_line = Side(style='thin', color='1B365D')
    double_bottom_line = Side(style='double', color='1B365D')

    border_data_cell = Border(left=thin_border_line, right=thin_border_line, top=thin_border_line, bottom=thin_border_line)
    border_total_row = Border(top=medium_top_line, bottom=double_bottom_line, left=thin_border_line, right=thin_border_line)

    # Standard clean numeric formats (without currency symbols)
    num_fmt = "#,##0.00"
    int_fmt = "#,##0"

    def create_sheet_banner(ws, title, subtitle, max_cols=10, fill_color=fill_banner_navy):
        ws.row_dimensions[1].height = 28
        ws.row_dimensions[2].height = 18
        ws.row_dimensions[3].height = 10
        
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_cols)
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max_cols)
        
        cell1 = ws.cell(row=1, column=1, value=f"  {title.upper()}")
        cell1.font = font_main_title
        cell1.fill = fill_color
        cell1.alignment = Alignment(horizontal="left", vertical="center")
        
        cell2 = ws.cell(row=2, column=1, value=f"  {subtitle} | Report Generated on: {date.today().strftime('%d-%B-%Y')}")
        cell2.font = font_sub_title
        cell2.fill = fill_color
        cell2.alignment = Alignment(horizontal="left", vertical="center")

        for r in range(1, 3):
            for c in range(1, max_cols + 1):
                ws.cell(row=r, column=c).fill = fill_color


    # =========================================================================
    # SHEET 1: SYSTEM CREDENTIALS & USER ACCESS (FIRST SHEET AS REQUESTED)
    # =========================================================================
    ws_users = wb.create_sheet(title="Credentials & Access")
    ws_users.views.sheetView[0].showGridLines = True
    create_sheet_banner(
        ws_users,
        "System Users, Roles & Login Credentials",
        "Master Directory of All Company Admins, Supplier Portal Users, Managers, Executives & End Users",
        max_cols=9,
        fill_color=fill_banner_dark
    )

    cred_headers = [
        ("User Category / Role", Alignment(horizontal="left")),
        ("Entity / Linked Organization", Alignment(horizontal="left")),
        ("Full Name / Display Name", Alignment(horizontal="left")),
        ("Username / Login ID", Alignment(horizontal="left")),
        ("Password", Alignment(horizontal="center")),
        ("Email Address", Alignment(horizontal="left")),
        ("Contact Phone", Alignment(horizontal="center")),
        ("Assigned Area / Location", Alignment(horizontal="left")),
        ("Account Status", Alignment(horizontal="center")),
    ]

    ws_users.row_dimensions[4].height = 26
    for col_idx, (col_name, align) in enumerate(cred_headers, start=1):
        cell = ws_users.cell(row=4, column=col_idx, value=col_name)
        cell.font = font_col_header
        cell.fill = fill_hdr_navy
        cell.alignment = Alignment(horizontal=align.horizontal, vertical="center", wrap_text=True)
        cell.border = border_data_cell

    ws_users.freeze_panes = "A5"

    row_num = 5

    # 1. Company Admin Users
    admin_users = Users.objects.select_related('companyid').all().order_by('userid')
    for u in admin_users:
        comp_name = u.companyid.companyname if u.companyid else "Super Admin"
        ws_users.row_dimensions[row_num].height = 20
        ws_users.cell(row=row_num, column=1, value="Company Admin").font = font_data_bold
        ws_users.cell(row=row_num, column=2, value=comp_name)
        ws_users.cell(row=row_num, column=3, value=u.username.capitalize()).font = font_data_bold
        ws_users.cell(row=row_num, column=4, value=u.username).font = font_data_code
        
        pwd_c = ws_users.cell(row=row_num, column=5, value=u.userpassword)
        pwd_c.alignment = Alignment(horizontal="center")
        pwd_c.font = font_data_bold
        pwd_c.fill = fill_badge_yellow

        ws_users.cell(row=row_num, column=6, value=u.useremail or '-')
        ws_users.cell(row=row_num, column=7, value='-').alignment = Alignment(horizontal="center")
        ws_users.cell(row=row_num, column=8, value=u.companyid.companylocation if u.companyid else '-')
        st_c = ws_users.cell(row=row_num, column=9, value="Active")
        st_c.alignment = Alignment(horizontal="center")
        st_c.fill = fill_badge_green

        for c in range(1, 10):
            ws_users.cell(row=row_num, column=c).border = border_data_cell
        row_num += 1

    # 2. Supplier Portal Users
    supp_users = Supplieruser.objects.select_related('supplierid', 'supplierid__companyid').all().order_by('supplieruserid')
    for su in supp_users:
        supp_name = su.supplierid.suppliername if su.supplierid else (su.suppliername or "Supplier")
        ws_users.row_dimensions[row_num].height = 20
        ws_users.cell(row=row_num, column=1, value="Supplier Admin").font = font_data_bold
        ws_users.cell(row=row_num, column=2, value=supp_name)
        ws_users.cell(row=row_num, column=3, value=su.suppliername or su.supplierusername).font = font_data_bold
        ws_users.cell(row=row_num, column=4, value=su.supplierusername).font = font_data_code
        
        pwd_c = ws_users.cell(row=row_num, column=5, value=su.supplieruserpassword)
        pwd_c.alignment = Alignment(horizontal="center")
        pwd_c.font = font_data_bold
        pwd_c.fill = fill_badge_yellow

        ws_users.cell(row=row_num, column=6, value=su.supplieruseremail or '-')
        ws_users.cell(row=row_num, column=7, value=su.supplieruserphone or '-').alignment = Alignment(horizontal="center")
        ws_users.cell(row=row_num, column=8, value=su.supplieruseraddress or '-')
        st_c = ws_users.cell(row=row_num, column=9, value="Active")
        st_c.alignment = Alignment(horizontal="center")
        st_c.fill = fill_badge_green

        for c in range(1, 10):
            ws_users.cell(row=row_num, column=c).border = border_data_cell
        row_num += 1

    # 3. Supplier Managers
    managers = SupplierManager.objects.select_related('supplier_user').all().order_by('manager_id')
    for m in managers:
        supp_name = m.supplier_user.suppliername if m.supplier_user else "Supplier Org"
        ws_users.row_dimensions[row_num].height = 20
        ws_users.cell(row=row_num, column=1, value="Supplier Area Manager").font = font_data_bold
        ws_users.cell(row=row_num, column=2, value=supp_name)
        ws_users.cell(row=row_num, column=3, value=m.manager_name).font = font_data_bold
        ws_users.cell(row=row_num, column=4, value=m.manager_username).font = font_data_code
        
        pwd_c = ws_users.cell(row=row_num, column=5, value=m.manager_password)
        pwd_c.alignment = Alignment(horizontal="center")
        pwd_c.font = font_data_bold
        pwd_c.fill = fill_badge_yellow

        ws_users.cell(row=row_num, column=6, value='-')
        ws_users.cell(row=row_num, column=7, value=m.manager_phone or '-').alignment = Alignment(horizontal="center")
        ws_users.cell(row=row_num, column=8, value=m.manager_area or '-')
        st_c = ws_users.cell(row=row_num, column=9, value="Active")
        st_c.alignment = Alignment(horizontal="center")
        st_c.fill = fill_badge_green

        for c in range(1, 10):
            ws_users.cell(row=row_num, column=c).border = border_data_cell
        row_num += 1

    # 4. Supplier Executives
    executives = SupplierExecutive.objects.select_related('manager', 'supplier_user').all().order_by('executiveid')
    for e in executives:
        supp_name = e.supplier_user.suppliername if e.supplier_user else (e.manager.manager_name if e.manager else "Supplier Org")
        ws_users.row_dimensions[row_num].height = 20
        ws_users.cell(row=row_num, column=1, value="Field Sales Executive").font = font_data_bold
        ws_users.cell(row=row_num, column=2, value=supp_name)
        ws_users.cell(row=row_num, column=3, value=e.executive_name).font = font_data_bold
        ws_users.cell(row=row_num, column=4, value=e.executive_username).font = font_data_code
        
        pwd_c = ws_users.cell(row=row_num, column=5, value=e.executive_password)
        pwd_c.alignment = Alignment(horizontal="center")
        pwd_c.font = font_data_bold
        pwd_c.fill = fill_badge_yellow

        ws_users.cell(row=row_num, column=6, value='-')
        ws_users.cell(row=row_num, column=7, value=e.executive_phone or '-').alignment = Alignment(horizontal="center")
        ws_users.cell(row=row_num, column=8, value=e.manager.manager_area if e.manager else '-')
        st_c = ws_users.cell(row=row_num, column=9, value="Active")
        st_c.alignment = Alignment(horizontal="center")
        st_c.fill = fill_badge_green

        for c in range(1, 10):
            ws_users.cell(row=row_num, column=c).border = border_data_cell
        row_num += 1

    # 5. Mobile App End Users / Customers
    end_users = EndUser.objects.all().order_by('endsuerid')
    for eu in end_users:
        ws_users.row_dimensions[row_num].height = 20
        ws_users.cell(row=row_num, column=1, value="Customer / App User").font = font_data_bold
        ws_users.cell(row=row_num, column=2, value="Retail Customer")
        ws_users.cell(row=row_num, column=3, value=eu.endusername).font = font_data_bold
        ws_users.cell(row=row_num, column=4, value=eu.endusername).font = font_data_code
        
        pwd_c = ws_users.cell(row=row_num, column=5, value=eu.enduserpassword or '-')
        pwd_c.alignment = Alignment(horizontal="center")
        pwd_c.font = font_data_bold
        pwd_c.fill = fill_badge_yellow

        ws_users.cell(row=row_num, column=6, value=eu.enduseremail or '-')
        ws_users.cell(row=row_num, column=7, value=eu.enduserphone or '-').alignment = Alignment(horizontal="center")
        ws_users.cell(row=row_num, column=8, value="Kerala, India")
        st_c = ws_users.cell(row=row_num, column=9, value="Active")
        st_c.alignment = Alignment(horizontal="center")
        st_c.fill = fill_badge_green

        for c in range(1, 10):
            ws_users.cell(row=row_num, column=c).border = border_data_cell
        row_num += 1

    ws_users.auto_filter.ref = f"A4:I{row_num - 1}"


    # =========================================================================
    # FETCH BILLS DATA FOR SHEETS 2-6
    # =========================================================================
    cursor = connection.cursor()

    cursor.execute("""
        SELECT cb.customerbillid, cb.customerid, c.customername, comp.companyid, comp.companyname,
               cb.customerbillno, cb.customerbilldate, cb.customerbillduedate, cb.customerbillamount,
               cb.customerbilltype, cb.narration, cb.cid
        FROM customerbill cb
        JOIN customer c ON cb.customerid = c.customerid
        JOIN company comp ON c.companyid = comp.companyid
        ORDER BY comp.companyid, cb.customerid, cb.customerbillid, cb.cid
    """)
    cust_bills_raw = cursor.fetchall()

    cursor.execute("""
        SELECT sb.supplierbillid, sb.supplierid, s.suppliername, comp.companyid, comp.companyname,
               sb.supplierbillno, sb.supplierbilldate, sb.supplierbillduedate, sb.supplierbillamount,
               sb.supplierbilltype, sb.narration, sb.sid
        FROM supplierbill sb
        JOIN supplier s ON sb.supplierid = s.supplierid
        JOIN company comp ON s.companyid = comp.companyid
        ORDER BY comp.companyid, sb.supplierid, sb.supplierbillid, sb.sid
    """)
    supp_bills_raw = cursor.fetchall()


    # =========================================================================
    # SHEET 2: EXECUTIVE FINANCIAL SUMMARY DASHBOARD
    # =========================================================================
    ws_exec = wb.create_sheet(title="Executive Summary")
    ws_exec.views.sheetView[0].showGridLines = True
    create_sheet_banner(
        ws_exec, 
        "Customer & Supplier Financial Executive Dashboard", 
        "High-Level Company Financial Summaries & Net Outstanding Balances", 
        max_cols=8,
        fill_color=fill_banner_navy
    )

    # 2A. Customer Overview by Company
    ws_exec.cell(row=4, column=1, value="1. CUSTOMER FINANCIAL OVERVIEW BY COMPANY").font = font_sec_heading
    ws_exec.row_dimensions[5].height = 24

    cust_summary_cols = [
        ("Company ID", Alignment(horizontal="center")),
        ("Company Name", Alignment(horizontal="left")),
        ("Total Customers", Alignment(horizontal="center")),
        ("Total Sales", Alignment(horizontal="right")),
        ("Total Receipts", Alignment(horizontal="right")),
        ("Credit Notes", Alignment(horizontal="right")),
        ("Net Outstanding", Alignment(horizontal="right")),
        ("Status", Alignment(horizontal="center")),
    ]

    for col_idx, (col_name, align) in enumerate(cust_summary_cols, start=1):
        cell = ws_exec.cell(row=5, column=col_idx, value=col_name)
        cell.font = font_col_header
        cell.fill = fill_hdr_blue
        cell.alignment = Alignment(horizontal=align.horizontal, vertical="center", wrap_text=True)
        cell.border = border_data_cell

    comp_cust = defaultdict(lambda: {'name': '', 'custs': set(), 'sales': Decimal(0), 'rec': Decimal(0), 'notes': Decimal(0)})
    for row in cust_bills_raw:
        comp_id, cust_id, comp_name = row[3], row[1], row[4]
        amt = Decimal(str(row[8] or 0))
        btype = (row[9] or '').strip().lower()
        comp_cust[comp_id]['name'] = comp_name
        comp_cust[comp_id]['custs'].add(cust_id)
        if any(m in btype for m in ['return', 'credit', 'note']):
            comp_cust[comp_id]['notes'] += amt
        elif 'receipt' in btype:
            comp_cust[comp_id]['rec'] += amt
        else:
            comp_cust[comp_id]['sales'] += amt

    row_num = 6
    tot_cust_c = 0
    tot_sales = Decimal(0)
    tot_rec = Decimal(0)
    tot_cnotes = Decimal(0)

    for comp_id in sorted(comp_cust.keys()):
        d = comp_cust[comp_id]
        c_count = len(d['custs'])
        s_val = d['sales']
        r_val = d['rec']
        n_val = d['notes']
        out_val = s_val - r_val - n_val

        tot_cust_c += c_count
        tot_sales += s_val
        tot_rec += r_val
        tot_cnotes += n_val

        ws_exec.row_dimensions[row_num].height = 20
        ws_exec.cell(row=row_num, column=1, value=f"COMP-{comp_id:02d}").alignment = Alignment(horizontal="center")
        ws_exec.cell(row=row_num, column=2, value=d['name']).font = font_data_bold
        ws_exec.cell(row=row_num, column=3, value=c_count).alignment = Alignment(horizontal="center")
        ws_exec.cell(row=row_num, column=4, value=float(s_val)).number_format = num_fmt
        ws_exec.cell(row=row_num, column=5, value=float(r_val)).number_format = num_fmt
        ws_exec.cell(row=row_num, column=6, value=float(n_val)).number_format = num_fmt
        
        out_c = ws_exec.cell(row=row_num, column=7, value=float(out_val))
        out_c.number_format = num_fmt
        out_c.font = font_data_bold
        out_c.fill = fill_badge_green if out_val >= 0 else fill_badge_red

        st_c = ws_exec.cell(row=row_num, column=8, value="Active Receivables" if out_val > 0 else "Settled")
        st_c.alignment = Alignment(horizontal="center")
        st_c.font = font_data

        for c in range(1, 9):
            ws_exec.cell(row=row_num, column=c).border = border_data_cell
        row_num += 1

    # Total Row
    ws_exec.row_dimensions[row_num].height = 22
    ws_exec.cell(row=row_num, column=1, value="GRAND TOTAL").alignment = Alignment(horizontal="center")
    ws_exec.cell(row=row_num, column=2, value=f"{len(comp_cust)} Companies Total").font = font_total_label
    ws_exec.cell(row=row_num, column=3, value=tot_cust_c).alignment = Alignment(horizontal="center")
    ws_exec.cell(row=row_num, column=4, value=float(tot_sales)).number_format = num_fmt
    ws_exec.cell(row=row_num, column=5, value=float(tot_rec)).number_format = num_fmt
    ws_exec.cell(row=row_num, column=6, value=float(tot_cnotes)).number_format = num_fmt
    tot_out_c = ws_exec.cell(row=row_num, column=7, value=float(tot_sales - tot_rec - tot_cnotes))
    tot_out_c.number_format = num_fmt
    ws_exec.cell(row=row_num, column=8, value="")

    for c in range(1, 9):
        cell = ws_exec.cell(row=row_num, column=c)
        cell.font = font_total_val
        cell.fill = fill_total_row
        cell.border = border_total_row

    # 2B. Supplier Overview by Company
    row_num += 3
    ws_exec.cell(row=row_num, column=1, value="2. SUPPLIER FINANCIAL OVERVIEW BY COMPANY").font = font_sec_heading
    row_num += 1
    ws_exec.row_dimensions[row_num].height = 24

    supp_summary_cols = [
        ("Company ID", Alignment(horizontal="center")),
        ("Company Name", Alignment(horizontal="left")),
        ("Total Suppliers", Alignment(horizontal="center")),
        ("Total Purchases", Alignment(horizontal="right")),
        ("Total Payments", Alignment(horizontal="right")),
        ("Debit Notes", Alignment(horizontal="right")),
        ("Net Outstanding", Alignment(horizontal="right")),
        ("Status", Alignment(horizontal="center")),
    ]

    for col_idx, (col_name, align) in enumerate(supp_summary_cols, start=1):
        cell = ws_exec.cell(row=row_num, column=col_idx, value=col_name)
        cell.font = font_col_header
        cell.fill = fill_hdr_darkgreen
        cell.alignment = Alignment(horizontal=align.horizontal, vertical="center", wrap_text=True)
        cell.border = border_data_cell

    row_num += 1

    comp_supp = defaultdict(lambda: {'name': '', 'supps': set(), 'pur': Decimal(0), 'pay': Decimal(0), 'notes': Decimal(0)})
    for row in supp_bills_raw:
        comp_id, supp_id, comp_name = row[3], row[1], row[4]
        amt = Decimal(str(row[8] or 0))
        btype = (row[9] or '').strip().lower()
        comp_supp[comp_id]['name'] = comp_name
        comp_supp[comp_id]['supps'].add(supp_id)
        if any(m in btype for m in ['return', 'debit', 'note']):
            comp_supp[comp_id]['notes'] += amt
        elif 'payment' in btype:
            comp_supp[comp_id]['pay'] += amt
        else:
            comp_supp[comp_id]['pur'] += amt

    tot_supp_c = 0
    tot_pur = Decimal(0)
    tot_pay = Decimal(0)
    tot_snotes = Decimal(0)

    for comp_id in sorted(comp_supp.keys()):
        d = comp_supp[comp_id]
        s_count = len(d['supps'])
        p_val = d['pur']
        pay_val = d['pay']
        n_val = d['notes']
        out_val = p_val - pay_val - n_val

        tot_supp_c += s_count
        tot_pur += p_val
        tot_pay += pay_val
        tot_snotes += n_val

        ws_exec.row_dimensions[row_num].height = 20
        ws_exec.cell(row=row_num, column=1, value=f"COMP-{comp_id:02d}").alignment = Alignment(horizontal="center")
        ws_exec.cell(row=row_num, column=2, value=d['name']).font = font_data_bold
        ws_exec.cell(row=row_num, column=3, value=s_count).alignment = Alignment(horizontal="center")
        ws_exec.cell(row=row_num, column=4, value=float(p_val)).number_format = num_fmt
        ws_exec.cell(row=row_num, column=5, value=float(pay_val)).number_format = num_fmt
        ws_exec.cell(row=row_num, column=6, value=float(n_val)).number_format = num_fmt
        
        out_c = ws_exec.cell(row=row_num, column=7, value=float(out_val))
        out_c.number_format = num_fmt
        out_c.font = font_data_bold
        out_c.fill = fill_badge_green if out_val >= 0 else fill_badge_red

        st_c = ws_exec.cell(row=row_num, column=8, value="Active Payables" if out_val > 0 else "Settled")
        st_c.alignment = Alignment(horizontal="center")
        st_c.font = font_data

        for c in range(1, 9):
            ws_exec.cell(row=row_num, column=c).border = border_data_cell
        row_num += 1

    # Total Row
    ws_exec.row_dimensions[row_num].height = 22
    ws_exec.cell(row=row_num, column=1, value="GRAND TOTAL").alignment = Alignment(horizontal="center")
    ws_exec.cell(row=row_num, column=2, value=f"{len(comp_supp)} Companies Total").font = font_total_label
    ws_exec.cell(row=row_num, column=3, value=tot_supp_c).alignment = Alignment(horizontal="center")
    ws_exec.cell(row=row_num, column=4, value=float(tot_pur)).number_format = num_fmt
    ws_exec.cell(row=row_num, column=5, value=float(tot_pay)).number_format = num_fmt
    ws_exec.cell(row=row_num, column=6, value=float(tot_snotes)).number_format = num_fmt
    tot_s_out_c = ws_exec.cell(row=row_num, column=7, value=float(tot_pur - tot_pay - tot_snotes))
    tot_s_out_c.number_format = num_fmt
    ws_exec.cell(row=row_num, column=8, value="")

    for c in range(1, 9):
        cell = ws_exec.cell(row=row_num, column=c)
        cell.font = font_total_val
        cell.fill = fill_total_row
        cell.border = border_total_row


    # =========================================================================
    # SHEET 3: CUSTOMER BILLS (DETAIL)
    # =========================================================================
    ws_cb_all = wb.create_sheet(title="Customer Bills (Detail)")
    ws_cb_all.views.sheetView[0].showGridLines = True
    create_sheet_banner(
        ws_cb_all,
        "Customer Billing Transactions & Ledger Details",
        "Itemized Record of Sales Invoices, Payment Receipts, and Credit/Return Notes",
        max_cols=11,
        fill_color=fill_banner_navy
    )

    cb_detail_headers = [
        ("Sl No", Alignment(horizontal="center")),
        ("Bill Group ID", Alignment(horizontal="center")),
        ("Company Name", Alignment(horizontal="left")),
        ("Customer ID", Alignment(horizontal="center")),
        ("Customer Name", Alignment(horizontal="left")),
        ("Bill / Invoice No", Alignment(horizontal="center")),
        ("Invoice Date", Alignment(horizontal="center")),
        ("Due Date", Alignment(horizontal="center")),
        ("Transaction Type", Alignment(horizontal="center")),
        ("Amount", Alignment(horizontal="right")),
        ("Narration / Description", Alignment(horizontal="left")),
    ]

    ws_cb_all.row_dimensions[4].height = 26
    for col_idx, (hdr_text, align) in enumerate(cb_detail_headers, start=1):
        cell = ws_cb_all.cell(row=4, column=col_idx, value=hdr_text)
        cell.font = font_col_header
        cell.fill = fill_hdr_blue
        cell.alignment = Alignment(horizontal=align.horizontal, vertical="center", wrap_text=True)
        cell.border = border_data_cell

    ws_cb_all.freeze_panes = "A5"
    ws_cb_all.auto_filter.ref = f"A4:K{len(cust_bills_raw) + 4}"

    row_num = 5
    tot_cb_amount = Decimal(0)
    for sl, row in enumerate(cust_bills_raw, start=1):
        bill_grp_id, cust_id, cust_name, comp_id, comp_name, bill_no, bdate, duedate, amt, btype, narr, cid = row
        amt_dec = Decimal(str(amt or 0))
        tot_cb_amount += amt_dec

        ws_cb_all.row_dimensions[row_num].height = 19
        ws_cb_all.cell(row=row_num, column=1, value=sl).alignment = Alignment(horizontal="center")
        ws_cb_all.cell(row=row_num, column=2, value=f"GRP-{bill_grp_id:03d}").alignment = Alignment(horizontal="center")
        ws_cb_all.cell(row=row_num, column=3, value=comp_name)
        ws_cb_all.cell(row=row_num, column=4, value=f"CUST-{cust_id:03d}").alignment = Alignment(horizontal="center")
        ws_cb_all.cell(row=row_num, column=5, value=cust_name).font = font_data_bold
        ws_cb_all.cell(row=row_num, column=6, value=bill_no).font = font_data_bold
        ws_cb_all.cell(row=row_num, column=7, value=str(bdate) if bdate else '').alignment = Alignment(horizontal="center")
        ws_cb_all.cell(row=row_num, column=8, value=str(duedate) if duedate else '-').alignment = Alignment(horizontal="center")
        
        type_c = ws_cb_all.cell(row=row_num, column=9, value=(btype or '').upper())
        type_c.alignment = Alignment(horizontal="center")
        if 'SALES' in (btype or '').upper():
            type_c.font = Font(name="Segoe UI", size=10, bold=True, color="1F4E79")
            type_c.fill = fill_badge_blue
        elif 'RECEIPT' in (btype or '').upper():
            type_c.font = Font(name="Segoe UI", size=10, bold=True, color="1E8449")
            type_c.fill = fill_badge_green
        elif 'RETURN' in (btype or '').upper() or 'CREDIT' in (btype or '').upper():
            type_c.font = Font(name="Segoe UI", size=10, bold=True, color="922B21")
            type_c.fill = fill_badge_red

        amt_c = ws_cb_all.cell(row=row_num, column=10, value=float(amt_dec))
        amt_c.number_format = num_fmt
        ws_cb_all.cell(row=row_num, column=11, value=narr or '')

        if row_num % 2 == 1:
            for c in [1, 2, 3, 4, 5, 6, 7, 8, 10, 11]:
                ws_cb_all.cell(row=row_num, column=c).fill = fill_zebra_light

        for c in range(1, 12):
            ws_cb_all.cell(row=row_num, column=c).border = border_data_cell
        row_num += 1

    # Total Row
    ws_cb_all.row_dimensions[row_num].height = 22
    ws_cb_all.cell(row=row_num, column=1, value="TOTAL").alignment = Alignment(horizontal="center")
    ws_cb_all.cell(row=row_num, column=2, value="")
    ws_cb_all.cell(row=row_num, column=3, value="")
    ws_cb_all.cell(row=row_num, column=4, value="")
    ws_cb_all.cell(row=row_num, column=5, value=f"{len(cust_bills_raw)} Transactions").font = font_total_label
    ws_cb_all.cell(row=row_num, column=6, value="")
    ws_cb_all.cell(row=row_num, column=7, value="")
    ws_cb_all.cell(row=row_num, column=8, value="")
    ws_cb_all.cell(row=row_num, column=9, value="")
    tot_amt_c = ws_cb_all.cell(row=row_num, column=10, value=float(tot_cb_amount))
    tot_amt_c.number_format = num_fmt
    ws_cb_all.cell(row=row_num, column=11, value="")

    for c in range(1, 12):
        cell = ws_cb_all.cell(row=row_num, column=c)
        cell.font = font_total_val
        cell.fill = fill_total_row
        cell.border = border_total_row


    # =============================================================
    # SHEET 4: CUSTOMER OUTSTANDING BALANCES (INDIVIDUAL BREAKDOWN)
    # =============================================================
    ws_cb_out = wb.create_sheet(title="Customer Outstanding")
    ws_cb_out.views.sheetView[0].showGridLines = True
    create_sheet_banner(
        ws_cb_out,
        "Customer-Wise Individual Outstanding Ledger Summary",
        "Separate Outstanding Balances, Total Sales, Receipts, and Returns per Customer",
        max_cols=9,
        fill_color=fill_banner_navy
    )

    cust_out_cols = [
        ("Sl No", Alignment(horizontal="center")),
        ("Customer ID", Alignment(horizontal="center")),
        ("Customer Name", Alignment(horizontal="left")),
        ("Company Name", Alignment(horizontal="left")),
        ("Total Invoices", Alignment(horizontal="center")),
        ("Total Sales", Alignment(horizontal="right")),
        ("Total Receipts", Alignment(horizontal="right")),
        ("Credit Notes", Alignment(horizontal="right")),
        ("Net Outstanding Balance", Alignment(horizontal="right")),
    ]

    ws_cb_out.row_dimensions[4].height = 26
    for col_idx, (col_name, align) in enumerate(cust_out_cols, start=1):
        cell = ws_cb_out.cell(row=4, column=col_idx, value=col_name)
        cell.font = font_col_header
        cell.fill = fill_hdr_teal
        cell.alignment = Alignment(horizontal=align.horizontal, vertical="center", wrap_text=True)
        cell.border = border_data_cell

    ws_cb_out.freeze_panes = "A5"

    cust_indiv = defaultdict(lambda: {'name': '', 'comp': '', 'sales': Decimal(0), 'rec': Decimal(0), 'notes': Decimal(0), 'cnt': 0})
    for row in cust_bills_raw:
        cust_id, cust_name, comp_name = row[1], row[2], row[4]
        amt = Decimal(str(row[8] or 0))
        btype = (row[9] or '').strip().lower()
        cust_indiv[cust_id]['name'] = cust_name
        cust_indiv[cust_id]['comp'] = comp_name
        cust_indiv[cust_id]['cnt'] += 1
        if any(m in btype for m in ['return', 'credit', 'note']):
            cust_indiv[cust_id]['notes'] += amt
        elif 'receipt' in btype:
            cust_indiv[cust_id]['rec'] += amt
        else:
            cust_indiv[cust_id]['sales'] += amt

    row_num = 5
    tot_ind_sales = Decimal(0)
    tot_ind_rec = Decimal(0)
    tot_ind_notes = Decimal(0)
    tot_ind_cnt = 0

    for sl, cust_id in enumerate(sorted(cust_indiv.keys()), start=1):
        d = cust_indiv[cust_id]
        s_val = d['sales']
        r_val = d['rec']
        n_val = d['notes']
        out_val = s_val - r_val - n_val

        tot_ind_sales += s_val
        tot_ind_rec += r_val
        tot_ind_notes += n_val
        tot_ind_cnt += d['cnt']

        ws_cb_out.row_dimensions[row_num].height = 20
        ws_cb_out.cell(row=row_num, column=1, value=sl).alignment = Alignment(horizontal="center")
        ws_cb_out.cell(row=row_num, column=2, value=f"CUST-{cust_id:03d}").alignment = Alignment(horizontal="center")
        ws_cb_out.cell(row=row_num, column=3, value=d['name']).font = font_data_bold
        ws_cb_out.cell(row=row_num, column=4, value=d['comp'])
        ws_cb_out.cell(row=row_num, column=5, value=d['cnt']).alignment = Alignment(horizontal="center")
        ws_cb_out.cell(row=row_num, column=6, value=float(s_val)).number_format = num_fmt
        ws_cb_out.cell(row=row_num, column=7, value=float(r_val)).number_format = num_fmt
        ws_cb_out.cell(row=row_num, column=8, value=float(n_val)).number_format = num_fmt
        
        out_c = ws_cb_out.cell(row=row_num, column=9, value=float(out_val))
        out_c.number_format = num_fmt
        out_c.font = font_data_bold
        out_c.fill = fill_badge_green if out_val >= 0 else fill_badge_red

        if row_num % 2 == 1:
            for c in range(1, 9):
                ws_cb_out.cell(row=row_num, column=c).fill = fill_zebra_light

        for c in range(1, 10):
            ws_cb_out.cell(row=row_num, column=c).border = border_data_cell
        row_num += 1

    # Grand Total
    ws_cb_out.row_dimensions[row_num].height = 22
    ws_cb_out.cell(row=row_num, column=1, value="GRAND TOTAL").alignment = Alignment(horizontal="center")
    ws_cb_out.cell(row=row_num, column=2, value="")
    ws_cb_out.cell(row=row_num, column=3, value=f"{len(cust_indiv)} Customers").font = font_total_label
    ws_cb_out.cell(row=row_num, column=4, value="")
    ws_cb_out.cell(row=row_num, column=5, value=tot_ind_cnt).alignment = Alignment(horizontal="center")
    ws_cb_out.cell(row=row_num, column=6, value=float(tot_ind_sales)).number_format = num_fmt
    ws_cb_out.cell(row=row_num, column=7, value=float(tot_ind_rec)).number_format = num_fmt
    ws_cb_out.cell(row=row_num, column=8, value=float(tot_ind_notes)).number_format = num_fmt
    tot_out_c = ws_cb_out.cell(row=row_num, column=9, value=float(tot_ind_sales - tot_ind_rec - tot_ind_notes))
    tot_out_c.number_format = num_fmt

    for c in range(1, 10):
        cell = ws_cb_out.cell(row=row_num, column=c)
        cell.font = font_total_val
        cell.fill = fill_total_row
        cell.border = border_total_row


    # =============================================================
    # SHEET 5: SUPPLIER BILLS (DETAIL)
    # =============================================================
    ws_sb_all = wb.create_sheet(title="Supplier Bills (Detail)")
    ws_sb_all.views.sheetView[0].showGridLines = True
    create_sheet_banner(
        ws_sb_all,
        "Supplier Purchase Invoices & Payment Ledger",
        "Itemized Record of Stock Purchases, Bank Payments, and Debit/Return Notes",
        max_cols=11,
        fill_color=fill_banner_navy
    )

    sb_detail_headers = [
        ("Sl No", Alignment(horizontal="center")),
        ("Bill Group ID", Alignment(horizontal="center")),
        ("Company Name", Alignment(horizontal="left")),
        ("Supplier ID", Alignment(horizontal="center")),
        ("Supplier Name", Alignment(horizontal="left")),
        ("Bill / Invoice No", Alignment(horizontal="center")),
        ("Invoice Date", Alignment(horizontal="center")),
        ("Due Date", Alignment(horizontal="center")),
        ("Transaction Type", Alignment(horizontal="center")),
        ("Amount", Alignment(horizontal="right")),
        ("Narration / Description", Alignment(horizontal="left")),
    ]

    ws_sb_all.row_dimensions[4].height = 26
    for col_idx, (hdr_text, align) in enumerate(sb_detail_headers, start=1):
        cell = ws_sb_all.cell(row=4, column=col_idx, value=hdr_text)
        cell.font = font_col_header
        cell.fill = fill_hdr_darkgreen
        cell.alignment = Alignment(horizontal=align.horizontal, vertical="center", wrap_text=True)
        cell.border = border_data_cell

    ws_sb_all.freeze_panes = "A5"
    ws_sb_all.auto_filter.ref = f"A4:K{len(supp_bills_raw) + 4}"

    row_num = 5
    tot_sb_amount = Decimal(0)
    for sl, row in enumerate(supp_bills_raw, start=1):
        bill_grp_id, supp_id, supp_name, comp_id, comp_name, bill_no, bdate, duedate, amt, btype, narr, sid = row
        amt_dec = Decimal(str(amt or 0))
        tot_sb_amount += amt_dec

        ws_sb_all.row_dimensions[row_num].height = 19
        ws_sb_all.cell(row=row_num, column=1, value=sl).alignment = Alignment(horizontal="center")
        ws_sb_all.cell(row=row_num, column=2, value=f"GRP-{bill_grp_id:03d}").alignment = Alignment(horizontal="center")
        ws_sb_all.cell(row=row_num, column=3, value=comp_name)
        ws_sb_all.cell(row=row_num, column=4, value=f"SUPP-{supp_id:03d}").alignment = Alignment(horizontal="center")
        ws_sb_all.cell(row=row_num, column=5, value=supp_name).font = font_data_bold
        ws_sb_all.cell(row=row_num, column=6, value=bill_no).font = font_data_bold
        ws_sb_all.cell(row=row_num, column=7, value=str(bdate) if bdate else '').alignment = Alignment(horizontal="center")
        ws_sb_all.cell(row=row_num, column=8, value=str(duedate) if duedate else '-').alignment = Alignment(horizontal="center")
        
        type_c = ws_sb_all.cell(row=row_num, column=9, value=(btype or '').upper())
        type_c.alignment = Alignment(horizontal="center")
        if 'PURCHASE' in (btype or '').upper():
            type_c.font = Font(name="Segoe UI", size=10, bold=True, color="5B2C6F")
            type_c.fill = fill_badge_blue
        elif 'PAYMENT' in (btype or '').upper():
            type_c.font = Font(name="Segoe UI", size=10, bold=True, color="1E8449")
            type_c.fill = fill_badge_green
        elif 'DEBIT' in (btype or '').upper() or 'RETURN' in (btype or '').upper():
            type_c.font = Font(name="Segoe UI", size=10, bold=True, color="922B21")
            type_c.fill = fill_badge_red

        amt_c = ws_sb_all.cell(row=row_num, column=10, value=float(amt_dec))
        amt_c.number_format = num_fmt
        ws_sb_all.cell(row=row_num, column=11, value=narr or '')

        if row_num % 2 == 1:
            for c in [1, 2, 3, 4, 5, 6, 7, 8, 10, 11]:
                ws_sb_all.cell(row=row_num, column=c).fill = fill_zebra_light

        for c in range(1, 12):
            ws_sb_all.cell(row=row_num, column=c).border = border_data_cell
        row_num += 1

    # Total Row
    ws_sb_all.row_dimensions[row_num].height = 22
    ws_sb_all.cell(row=row_num, column=1, value="TOTAL").alignment = Alignment(horizontal="center")
    ws_sb_all.cell(row=row_num, column=2, value="")
    ws_sb_all.cell(row=row_num, column=3, value="")
    ws_sb_all.cell(row=row_num, column=4, value="")
    ws_sb_all.cell(row=row_num, column=5, value=f"{len(supp_bills_raw)} Transactions").font = font_total_label
    ws_sb_all.cell(row=row_num, column=6, value="")
    ws_sb_all.cell(row=row_num, column=7, value="")
    ws_sb_all.cell(row=row_num, column=8, value="")
    ws_sb_all.cell(row=row_num, column=9, value="")
    tot_amt_c = ws_sb_all.cell(row=row_num, column=10, value=float(tot_sb_amount))
    tot_amt_c.number_format = num_fmt
    ws_sb_all.cell(row=row_num, column=11, value="")

    for c in range(1, 12):
        cell = ws_sb_all.cell(row=row_num, column=c)
        cell.font = font_total_val
        cell.fill = fill_total_row
        cell.border = border_total_row


    # =============================================================
    # SHEET 6: SUPPLIER OUTSTANDING BALANCES (INDIVIDUAL BREAKDOWN)
    # =============================================================
    ws_sb_out = wb.create_sheet(title="Supplier Outstanding")
    ws_sb_out.views.sheetView[0].showGridLines = True
    create_sheet_banner(
        ws_sb_out,
        "Supplier-Wise Individual Outstanding Ledger Summary",
        "Separate Outstanding Balances, Total Purchases, Payments, and Debit Notes per Supplier",
        max_cols=9,
        fill_color=fill_banner_navy
    )

    supp_out_cols = [
        ("Sl No", Alignment(horizontal="center")),
        ("Supplier ID", Alignment(horizontal="center")),
        ("Supplier Name", Alignment(horizontal="left")),
        ("Company Name", Alignment(horizontal="left")),
        ("Total Invoices", Alignment(horizontal="center")),
        ("Total Purchases", Alignment(horizontal="right")),
        ("Total Payments", Alignment(horizontal="right")),
        ("Debit Notes", Alignment(horizontal="right")),
        ("Net Outstanding Payable", Alignment(horizontal="right")),
    ]

    ws_sb_out.row_dimensions[4].height = 26
    for col_idx, (col_name, align) in enumerate(supp_out_cols, start=1):
        cell = ws_sb_out.cell(row=4, column=col_idx, value=col_name)
        cell.font = font_col_header
        cell.fill = fill_hdr_purple
        cell.alignment = Alignment(horizontal=align.horizontal, vertical="center", wrap_text=True)
        cell.border = border_data_cell

    ws_sb_out.freeze_panes = "A5"

    supp_indiv = defaultdict(lambda: {'name': '', 'comp': '', 'pur': Decimal(0), 'pay': Decimal(0), 'notes': Decimal(0), 'cnt': 0})
    for row in supp_bills_raw:
        supp_id, supp_name, comp_name = row[1], row[2], row[4]
        amt = Decimal(str(row[8] or 0))
        btype = (row[9] or '').strip().lower()
        supp_indiv[supp_id]['name'] = supp_name
        supp_indiv[supp_id]['comp'] = comp_name
        supp_indiv[supp_id]['cnt'] += 1
        if any(m in btype for m in ['return', 'debit', 'note']):
            supp_indiv[supp_id]['notes'] += amt
        elif 'payment' in btype:
            supp_indiv[supp_id]['pay'] += amt
        else:
            supp_indiv[supp_id]['pur'] += amt

    row_num = 5
    tot_ind_pur = Decimal(0)
    tot_ind_pay = Decimal(0)
    tot_ind_snotes = Decimal(0)
    tot_ind_scnt = 0

    for sl, supp_id in enumerate(sorted(supp_indiv.keys()), start=1):
        d = supp_indiv[supp_id]
        p_val = d['pur']
        pay_val = d['pay']
        n_val = d['notes']
        out_val = p_val - pay_val - n_val

        tot_ind_pur += p_val
        tot_ind_pay += pay_val
        tot_ind_snotes += n_val
        tot_ind_scnt += d['cnt']

        ws_sb_out.row_dimensions[row_num].height = 20
        ws_sb_out.cell(row=row_num, column=1, value=sl).alignment = Alignment(horizontal="center")
        ws_sb_out.cell(row=row_num, column=2, value=f"SUPP-{supp_id:03d}").alignment = Alignment(horizontal="center")
        ws_sb_out.cell(row=row_num, column=3, value=d['name']).font = font_data_bold
        ws_sb_out.cell(row=row_num, column=4, value=d['comp'])
        ws_sb_out.cell(row=row_num, column=5, value=d['cnt']).alignment = Alignment(horizontal="center")
        ws_sb_out.cell(row=row_num, column=6, value=float(p_val)).number_format = num_fmt
        ws_sb_out.cell(row=row_num, column=7, value=float(pay_val)).number_format = num_fmt
        ws_sb_out.cell(row=row_num, column=8, value=float(n_val)).number_format = num_fmt
        
        out_c = ws_sb_out.cell(row=row_num, column=9, value=float(out_val))
        out_c.number_format = num_fmt
        out_c.font = font_data_bold
        out_c.fill = fill_badge_green if out_val >= 0 else fill_badge_red

        if row_num % 2 == 1:
            for c in range(1, 9):
                ws_sb_out.cell(row=row_num, column=c).fill = fill_zebra_light

        for c in range(1, 10):
            ws_sb_out.cell(row=row_num, column=c).border = border_data_cell
        row_num += 1

    # Grand Total
    ws_sb_out.row_dimensions[row_num].height = 22
    ws_sb_out.cell(row=row_num, column=1, value="GRAND TOTAL").alignment = Alignment(horizontal="center")
    ws_sb_out.cell(row=row_num, column=2, value="")
    ws_sb_out.cell(row=row_num, column=3, value=f"{len(supp_indiv)} Suppliers").font = font_total_label
    ws_sb_out.cell(row=row_num, column=4, value="")
    ws_sb_out.cell(row=row_num, column=5, value=tot_ind_scnt).alignment = Alignment(horizontal="center")
    ws_sb_out.cell(row=row_num, column=6, value=float(tot_ind_pur)).number_format = num_fmt
    ws_sb_out.cell(row=row_num, column=7, value=float(tot_ind_pay)).number_format = num_fmt
    ws_sb_out.cell(row=row_num, column=8, value=float(tot_ind_snotes)).number_format = num_fmt
    tot_s_out_c = ws_sb_out.cell(row=row_num, column=9, value=float(tot_ind_pur - tot_ind_pay - tot_ind_snotes))
    tot_s_out_c.number_format = num_fmt

    for c in range(1, 10):
        cell = ws_sb_out.cell(row=row_num, column=c)
        cell.font = font_total_val
        cell.fill = fill_total_row
        cell.border = border_total_row


    # =============================================================
    # SHEET 7: MASTER PRODUCTS CATALOG (PRODUCTS SHEET AS REQUESTED)
    # =============================================================
    ws_prod = wb.create_sheet(title="Products Catalog")
    ws_prod.views.sheetView[0].showGridLines = True
    create_sheet_banner(
        ws_prod,
        "Master Products Catalog",
        "Complete Directory of Products, Categories, Prices, and Stock Units per Company",
        max_cols=9,
        fill_color=fill_banner_navy
    )

    prod_headers = [
        ("Sl No", Alignment(horizontal="center")),
        ("Product ID", Alignment(horizontal="center")),
        ("Product Name", Alignment(horizontal="left")),
        ("Category", Alignment(horizontal="left")),
        ("Unit", Alignment(horizontal="center")),
        ("Selling Price", Alignment(horizontal="right")),
        ("Company Name", Alignment(horizontal="left")),
        ("Date Added", Alignment(horizontal="center")),
        ("Add Type", Alignment(horizontal="center")),
    ]

    ws_prod.row_dimensions[4].height = 26
    for col_idx, (hdr_text, align) in enumerate(prod_headers, start=1):
        cell = ws_prod.cell(row=4, column=col_idx, value=hdr_text)
        cell.font = font_col_header
        cell.fill = fill_hdr_darkamber
        cell.alignment = Alignment(horizontal=align.horizontal, vertical="center", wrap_text=True)
        cell.border = border_data_cell

    ws_prod.freeze_panes = "A5"

    products = Products.objects.select_related('companyid', 'productcategoryid', 'productunitid').all().order_by('companyid', 'productid')
    ws_prod.auto_filter.ref = f"A4:I{len(products) + 4}"

    row_num = 5
    tot_prod_price = Decimal(0)
    for sl, p in enumerate(products, start=1):
        comp_name = p.companyid.companyname if p.companyid else "Global"
        cat_name = p.productcategoryid.productcategoryname if p.productcategoryid else "General"
        unit_name = p.productunitid.productunitname if p.productunitid else "Nos"
        price = Decimal(str(p.productprice or 0))
        tot_prod_price += price

        ws_prod.row_dimensions[row_num].height = 20
        ws_prod.cell(row=row_num, column=1, value=sl).alignment = Alignment(horizontal="center")
        ws_prod.cell(row=row_num, column=2, value=f"PROD-{p.productid:03d}").alignment = Alignment(horizontal="center")
        ws_prod.cell(row=row_num, column=3, value=p.productname).font = font_data_bold
        ws_prod.cell(row=row_num, column=4, value=cat_name)
        ws_prod.cell(row=row_num, column=5, value=unit_name).alignment = Alignment(horizontal="center")
        
        pr_cell = ws_prod.cell(row=row_num, column=6, value=float(price))
        pr_cell.number_format = num_fmt
        pr_cell.font = font_data_bold

        ws_prod.cell(row=row_num, column=7, value=comp_name)
        ws_prod.cell(row=row_num, column=8, value=str(p.dateadded) if p.dateadded else '-').alignment = Alignment(horizontal="center")
        ws_prod.cell(row=row_num, column=9, value=p.addtype or 'Single').alignment = Alignment(horizontal="center")

        if row_num % 2 == 1:
            for c in range(1, 10):
                ws_prod.cell(row=row_num, column=c).fill = fill_zebra_light

        for c in range(1, 10):
            ws_prod.cell(row=row_num, column=c).border = border_data_cell
        row_num += 1

    # Total Row
    ws_prod.row_dimensions[row_num].height = 22
    ws_prod.cell(row=row_num, column=1, value="TOTAL").alignment = Alignment(horizontal="center")
    ws_prod.cell(row=row_num, column=2, value="")
    ws_prod.cell(row=row_num, column=3, value=f"{len(products)} Products Cataloged").font = font_total_label
    ws_prod.cell(row=row_num, column=4, value="")
    ws_prod.cell(row=row_num, column=5, value="")
    tot_pr_c = ws_prod.cell(row=row_num, column=6, value=float(tot_prod_price))
    tot_pr_c.number_format = num_fmt
    ws_prod.cell(row=row_num, column=7, value="")
    ws_prod.cell(row=row_num, column=8, value="")
    ws_prod.cell(row=row_num, column=9, value="")

    for c in range(1, 10):
        cell = ws_prod.cell(row=row_num, column=c)
        cell.font = font_total_val
        cell.fill = fill_total_row
        cell.border = border_total_row


    # =============================================================
    # SHEET 8: SUPPLIER PRODUCTS & PRICING (FINAL SHEET AS REQUESTED)
    # =============================================================
    ws_sp = wb.create_sheet(title="Supplier Products")
    ws_sp.views.sheetView[0].showGridLines = True
    create_sheet_banner(
        ws_sp,
        "Supplier Products & Wholesale Price Offerings",
        "Mapping of Supplier Wholesale Prices, Catalog Prices, Profit Margins, and Active Procurement Status",
        max_cols=10,
        fill_color=fill_banner_navy
    )

    sp_headers = [
        ("Sl No", Alignment(horizontal="center")),
        ("Mapping ID", Alignment(horizontal="center")),
        ("Company Name", Alignment(horizontal="left")),
        ("Supplier ID & Name", Alignment(horizontal="left")),
        ("Product ID & Name", Alignment(horizontal="left")),
        ("Supplier Wholesale Price", Alignment(horizontal="right")),
        ("Retail / Catalog Price", Alignment(horizontal="right")),
        ("Retail Margin", Alignment(horizontal="right")),
        ("Margin %", Alignment(horizontal="right")),
        ("Procurement Status", Alignment(horizontal="center")),
    ]

    ws_sp.row_dimensions[4].height = 26
    for col_idx, (hdr_text, align) in enumerate(sp_headers, start=1):
        cell = ws_sp.cell(row=4, column=col_idx, value=hdr_text)
        cell.font = font_col_header
        cell.fill = fill_hdr_steel
        cell.alignment = Alignment(horizontal=align.horizontal, vertical="center", wrap_text=True)
        cell.border = border_data_cell

    ws_sp.freeze_panes = "A5"

    supp_prods = SupplierProduct.objects.select_related('company', 'supplier', 'product').all().order_by('company', 'supplier', 'product')
    ws_sp.auto_filter.ref = f"A4:J{len(supp_prods) + 4}"

    row_num = 5
    tot_sp_price = Decimal(0)
    tot_cat_price = Decimal(0)

    for sl, sp in enumerate(supp_prods, start=1):
        comp_name = sp.company.companyname if sp.company else "Company"
        supp_disp = f"SUPP-{sp.supplier.supplierid:03d} | {sp.supplier.suppliername}"
        prod_disp = f"PROD-{sp.product.productid:03d} | {sp.product.productname}"
        
        s_price = Decimal(str(sp.supplier_price or 0)) if sp.supplier_price is not None else Decimal(0)
        c_price = Decimal(str(sp.product.productprice or 0))
        margin = c_price - s_price if s_price > 0 else Decimal(0)
        margin_pct = (margin / c_price * 100) if c_price > 0 and s_price > 0 else Decimal(0)

        tot_sp_price += s_price
        tot_cat_price += c_price

        ws_sp.row_dimensions[row_num].height = 20
        ws_sp.cell(row=row_num, column=1, value=sl).alignment = Alignment(horizontal="center")
        ws_sp.cell(row=row_num, column=2, value=f"MAP-{sp.id:03d}").alignment = Alignment(horizontal="center")
        ws_sp.cell(row=row_num, column=3, value=comp_name)
        ws_sp.cell(row=row_num, column=4, value=supp_disp).font = font_data_bold
        ws_sp.cell(row=row_num, column=5, value=prod_disp).font = font_data_bold
        
        sp_c = ws_sp.cell(row=row_num, column=6, value=float(s_price) if s_price > 0 else '-')
        if s_price > 0:
            sp_c.number_format = num_fmt
            sp_c.font = font_data_bold
        else:
            sp_c.alignment = Alignment(horizontal="center")

        cp_c = ws_sp.cell(row=row_num, column=7, value=float(c_price))
        cp_c.number_format = num_fmt
        cp_c.font = font_data_bold

        mg_c = ws_sp.cell(row=row_num, column=8, value=float(margin) if s_price > 0 else '-')
        if s_price > 0:
            mg_c.number_format = num_fmt
            mg_c.fill = fill_badge_green if margin > 0 else fill_badge_red
        else:
            mg_c.alignment = Alignment(horizontal="center")

        pct_c = ws_sp.cell(row=row_num, column=9, value=f"{margin_pct:.1f}%" if s_price > 0 else '-')
        pct_c.alignment = Alignment(horizontal="right")

        st_badge = ws_sp.cell(row=row_num, column=10, value="Active Connected" if sp.is_active else "Inactive")
        st_badge.alignment = Alignment(horizontal="center")
        st_badge.fill = fill_badge_green if sp.is_active else fill_badge_red

        if row_num % 2 == 1:
            for c in range(1, 11):
                if c not in [8, 10]:
                    ws_sp.cell(row=row_num, column=c).fill = fill_zebra_light

        for c in range(1, 11):
            ws_sp.cell(row=row_num, column=c).border = border_data_cell
        row_num += 1

    # Total Row
    ws_sp.row_dimensions[row_num].height = 22
    ws_sp.cell(row=row_num, column=1, value="TOTAL").alignment = Alignment(horizontal="center")
    ws_sp.cell(row=row_num, column=2, value="")
    ws_sp.cell(row=row_num, column=3, value=f"{len(supp_prods)} Product Mappings").font = font_total_label
    ws_sp.cell(row=row_num, column=4, value="")
    ws_sp.cell(row=row_num, column=5, value="")
    tot_sp_c = ws_sp.cell(row=row_num, column=6, value=float(tot_sp_price))
    tot_sp_c.number_format = num_fmt
    tot_cp_c = ws_sp.cell(row=row_num, column=7, value=float(tot_cat_price))
    tot_cp_c.number_format = num_fmt
    tot_mg_c = ws_sp.cell(row=row_num, column=8, value=float(tot_cat_price - tot_sp_price))
    tot_mg_c.number_format = num_fmt
    ws_sp.cell(row=row_num, column=9, value="")
    ws_sp.cell(row=row_num, column=10, value="")

    for c in range(1, 11):
        cell = ws_sp.cell(row=row_num, column=c)
        cell.font = font_total_val
        cell.fill = fill_total_row
        cell.border = border_total_row


    # =============================================================
    # AUTO-FIT COLUMN WIDTHS ACROSS ALL 8 SHEETS
    # =============================================================
    for sheet in wb.worksheets:
        for col in sheet.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.row in [1, 2, 3]:
                    continue
                val_str = str(cell.value or '')
                if cell.number_format == num_fmt and isinstance(cell.value, (int, float)):
                    val_str = f"{cell.value:,.2f}"
                max_len = max(max_len, len(val_str))
            sheet.column_dimensions[col_letter].width = max(max_len + 4, 13)

    abs_output_path = os.path.abspath(output_path)
    wb.save(abs_output_path)
    print(f"\n--> Full Enterprise Multi-Sheet Excel Workbook Successfully Generated at:")
    print(f"    {abs_output_path}")
    print("=" * 75)
    return abs_output_path

if __name__ == '__main__':
    export_bills_to_excel()
